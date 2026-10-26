import openpyxl
import os

files = [
    'mass_update_basic_info_381421545_20260911100357.xlsx',
    'mass_update_sales_info_381421545_20260911100617.xlsx',
    'mass_update_shipping_info_381421545_20260911100313.xlsx',
    'mass_update_media_info_381421545_20260911100714.xlsx'
]

for f in files:
    if os.path.exists(f):
        wb = openpyxl.load_workbook(f, data_only=True)
        print('========================================')
        print('FILE:', f, wb.sheetnames)
        for s in wb.sheetnames:
            ws = wb[s]
            print('  Sheet:', s, 'Max rows:', ws.max_row, 'Max cols:', ws.max_column)
            for i in range(1, min(15, ws.max_row + 1)):
                row_vals = [ws.cell(row=i, column=j).value for j in range(1, min(12, ws.max_column + 1))]
                # print non-empty rows
                if any(row_vals):
                    print(f'    Row {i:2d}:', [str(v)[:20] if v is not None else '' for v in row_vals])
